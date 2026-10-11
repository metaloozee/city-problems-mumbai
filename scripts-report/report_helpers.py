#!/usr/bin/env python3
"""Build NLP Mini Project Report DOCX matching SLRTCE college template.

Usage (from repo root):
  .venv-report/bin/python scripts-report/build_report.py
  libreoffice --headless --convert-to pdf --outdir report/ report/NLP_Mini_Project_Report.docx
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Cm, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
FIGS = REPORT / "figures"
OUT_DOCX = REPORT / "NLP_Mini_Project_Report.docx"

FONT = "Liberation Serif"
FONT_MONO = "Liberation Mono"


def set_run_font(run, size_pt=12, bold=False, italic=False, font=FONT):
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size_pt)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_para(
    doc,
    text="",
    *,
    size=12,
    bold=False,
    italic=False,
    align="left",
    space_before=0,
    space_after=6,
    line_spacing=1.15,
    font=FONT,
    page_break_before=False,
):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if page_break_before:
        pf.page_break_before = True

    if text:
        run = p.add_run(text)
        set_run_font(run, size_pt=size, bold=bold, italic=italic, font=font)
    return p


def add_mixed_para(doc, parts, *, align="justify", size=12, space_before=0, space_after=10, line_spacing=1.15):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    for text, bold in parts:
        run = p.add_run(text)
        set_run_font(run, size_pt=size, bold=bold)
    return p


def add_code_block(doc, code: str, size=9):
    for line in code.splitlines():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.4)
        run = p.add_run(line if line else " ")
        set_run_font(run, size_pt=size, bold=False, font=FONT_MONO)


def section_heading(doc, text):
    return add_para(
        doc,
        text,
        size=14,
        bold=True,
        align="center",
        space_after=14,
        line_spacing=1.15,
        page_break_before=True,
    )


def build():
    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ---------- COVER ----------
    add_para(doc, "", space_after=12)
    add_para(doc, "NLP Mini Project Report", size=16, bold=True, align="center", space_after=16, line_spacing=1.0)
    add_para(doc, "Fake News Detection using NLP and ML", size=20, bold=True, align="center", space_before=4, space_after=16, line_spacing=1.15)
    add_para(doc, "Submitted fulfilment of the requirements of the", size=12, align="center", space_after=0, line_spacing=1.15)
    add_para(doc, "NLP Mini Project Semester VII in degree of", size=12, align="center", space_after=0, line_spacing=1.15)
    add_para(doc, "Bachelor of Computer Engineering", size=12, align="center", space_after=12, line_spacing=1.15)
    add_para(doc, "BY", size=12, align="center", space_after=6, line_spacing=1.0)
    for name in ["Ayan Parkar (63)", "Sourabh Patel (64)", "Mohammed Owaiz Shaikh (65)"]:
        add_para(doc, name, size=13, bold=True, align="center", space_after=2, line_spacing=1.15)
    add_para(doc, "", space_after=8)
    add_para(doc, "Guide", size=12, align="center", space_after=2, line_spacing=1.0)
    add_para(doc, "Asst. Prof. Manasi Churi", size=12, align="center", space_after=12, line_spacing=1.15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(FIGS / "college-logo.png"), width=Inches(1.55))

    add_para(doc, "Department Of Computer Engineering", size=13, bold=True, align="center", space_after=2, line_spacing=1.15)
    add_para(doc, "Shree L. R. Tiwari College Of Engineering", size=13, bold=True, align="center", space_after=2, line_spacing=1.15)
    add_para(doc, "Kanakia Park, Mira Road (E), Thane -401 107, Maharashtra.", size=12, bold=True, align="center", space_after=12, line_spacing=1.15)
    add_para(doc, "2026-27", size=14, bold=True, align="center", space_after=0, line_spacing=1.0)

    # ---------- APPROVAL ----------
    add_para(doc, "NLP Mini Project Report Approval", size=16, bold=True, align="center", space_after=28, line_spacing=1.15, page_break_before=True)
    add_mixed_para(
        doc,
        [
            ("This project report entitled ", False),
            ("\u201cFake News Detection using NLP and ML\u201d", True),
            (" by ", False),
            ("Mr. Ayan Parkar (63), Mr. Sourabh Patel (64) and Mr. Mohammed Owaiz Shaikh (65)", True),
            (
                " is approved for the Mini Project Semester VII project work of B.E CMPN at Shree L.R.Tiwari College of Engineering, in the academic year 2026-2027.",
                False,
            ),
        ],
        align="justify", size=12, space_after=6, line_spacing=1.5,
    )
    for _ in range(10):
        add_para(doc, "", space_after=12)
    add_para(doc, "Internal Examiner", size=12, bold=True, align="right", space_after=0)

    # ---------- DECLARATION ----------
    add_para(doc, "Declaration", size=16, bold=True, align="center", space_after=20, line_spacing=1.15, page_break_before=True)
    decl = (
        "I declare that this written submission represents my ideas in my own words and where "
        "others' ideas or words have been included, I have adequately cited and referenced the "
        "original sources. I also declare that I have adhered to all principles of academic "
        "honesty and integrity and have not misrepresented or fabricated or falsified any "
        "idea/data/fact/source in my submission. I understand that any violation of the above "
        "will be cause for disciplinary action by the Institute and can also evoke penal action "
        "from the sources which have thus not been properly cited or from whom proper permission "
        "has not been taken when needed."
    )
    add_para(doc, decl, size=12, align="justify", space_after=36, line_spacing=1.5)
    for name in ["Mr. Ayan Parkar (63)", "Mr. Sourabh Patel (64)", "Mr. Mohammed Owaiz Shaikh (65)"]:
        add_para(doc, "________________________", size=12, align="right", space_before=18, space_after=0, line_spacing=1.0)
        add_para(doc, "Signature", size=11, align="right", space_after=0, line_spacing=1.0)
        add_para(doc, name, size=12, bold=True, align="right", space_after=6, line_spacing=1.15)
    add_para(doc, "", space_after=36)
    add_para(doc, "Date:", size=12, align="left", space_after=6)
    add_para(doc, "Place:", size=12, align="left", space_after=0)

    # ---------- TOC ----------
    add_para(doc, "Table of Contents", size=14, bold=True, align="left", space_after=14, line_spacing=1.15, page_break_before=True)
    for item in [
        "1. Introduction",
        "2. System Architecture",
        "3. Implementation",
        "4. Results",
        "5. Conclusion",
        "6. References",
    ]:
        add_para(doc, item, size=12, align="left", space_after=4, line_spacing=1.5)

    # ---------- 1. INTRODUCTION ----------
    section_heading(doc, "1. Introduction")
    for t in [
        (
            "Online news spreads quickly across social media and websites. Alongside legitimate "
            "reporting, fabricated or misleading articles can reach large audiences before any "
            "manual fact-check is possible. Automatically ranking a short news text as likely "
            "fake or real is a classic Natural Language Processing (NLP) and machine learning "
            "problem: the input is free-form English, and the output is a binary label with a "
            "confidence score."
        ),
        (
            "Classical bag-of-words models remain a strong baseline for this task. Term "
            "Frequency\u2013Inverse Document Frequency (TF-IDF) converts an article into a sparse "
            "numeric vector that down-weights common words and up-weights distinctive terms. A "
            "logistic regression classifier then learns a linear decision boundary between the "
            "FAKE and REAL classes. The approach is transparent, fast to train, and easy to "
            "export for browser-side inference without a Python runtime."
        ),
        (
            "This mini-project implements Fake News Detection using NLP and ML as a laboratory "
            "demo for the NLP Mini Project (Semester VII). A scikit-learn pipeline (TF-IDF + "
            "logistic regression) is trained offline on the public McIntire fake/real news "
            "corpus. Model parameters are exported to JSON. A Next.js web application scores "
            "pasted articles in TypeScript and shows FAKE or REAL with class probabilities. "
            "Hold-out accuracy and F1 scores are about 0.92 on a stratified 20% test split."
        ),
    ]:
        add_para(doc, t, size=12, align="justify", space_after=10, line_spacing=1.5)

    add_para(doc, "Objectives", size=12, bold=True, align="left", space_before=6, space_after=6)
    for i, o in enumerate([
        "Train a TF-IDF + logistic regression classifier on a public fake/real news dataset.",
        "Evaluate hold-out accuracy, per-class F1, and the confusion matrix.",
        "Export the fitted vocabulary, IDF weights, and coefficients for pure JavaScript inference.",
        "Build a Next.js UI and API so users can paste text and see FAKE/REAL predictions.",
        "Document architecture, implementation, and results for academic evaluation.",
    ], 1):
        add_para(doc, f"{i}. {o}", size=12, align="justify", space_after=4, line_spacing=1.35)

    add_para(doc, "Scope (MVP)", size=12, bold=True, align="left", space_before=10, space_after=6)
    for s in [
        "Binary FAKE/REAL classification on English news text; not a general fact-checker.",
        "McIntire corpus (mostly 2016 US politics); domain shift is expected outside that set.",
        "No transformer models; classical TF-IDF + logistic regression only.",
        "Browser/demo inference uses exported JSON; Python is required only for (re)training.",
    ]:
        add_para(doc, f"\u2022  {s}", size=12, align="justify", space_after=4, line_spacing=1.35)

    # ---------- 2. SYSTEM ARCHITECTURE ----------
    section_heading(doc, "2. System Architecture")
    add_para(
        doc,
        (
            "The system has two roles. A developer (offline) trains and exports the model. An "
            "end user opens the web demo, pastes a headline and article body (or loads a hold-out "
            "example), and receives a FAKE or REAL label with confidence and class probabilities."
        ),
        size=12, align="justify", space_after=10, line_spacing=1.5,
    )
    add_para(doc, "The system is organized in three layers:", size=12, align="justify", space_after=6, line_spacing=1.5)
    for bold_part, rest in [
        ("Training layer \u2014 ", "ml/train.py downloads the CSV if needed, fits TF-IDF (max_features=5000, English stop words) on the training split only, fits logistic regression (liblinear), evaluates the hold-out set, and writes models/metrics.json, models/model.joblib, web/lib/model.json, and web/lib/samples.json."),
        ("Inference layer \u2014 ", "web/lib/predict.ts reimplements tokenize \u2192 TF-IDF lookup \u2192 L2 normalize \u2192 sigmoid using the exported vocabulary, IDF, coefficients, and intercept. The Next.js route POST /api/predict returns label and probabilities."),
        ("Presentation layer \u2014 ", "A Next.js App Router page (web/app/page.tsx) provides a textarea, Load fake/real example buttons, Predict, and a result panel showing FAKE or REAL with a confidence bar."),
    ]:
        add_mixed_para(doc, [(f"\u2022  {bold_part}", True), (rest, False)], align="justify", size=12, space_after=6, line_spacing=1.4)

    add_para(doc, "Logical architecture:", size=12, bold=True, align="left", space_before=8, space_after=6)
    arch = """+----------------------+     train / export      +---------------------------+
|  ml/train.py         | ----------------------> |  web/lib/model.json       |
|  (sklearn TF-IDF +   |                         |  vocab, idf, coef, b      |
|   LogisticRegression)|                         +---------------------------+
+----------------------+                                    |
                                                            v
+----------------------+     POST { text }       +---------------------------+
|  Next.js UI          | <---------------------> |  /api/predict + predict.ts|
|  (paste / examples)  |     label + probs       |  (pure JS TF-IDF+sigmoid) |
+----------------------+                         +---------------------------+"""
    add_code_block(doc, arch, size=8)

    add_para(doc, "Data flow and scoring", size=12, bold=True, align="left", space_before=12, space_after=6)
    add_para(
        doc,
        (
            "Each document is built as title + blank line + body. At request time the text is "
            "lowercased and tokenized with a pattern matching sklearn\u2019s \\b\\w\\w+\\b on ASCII. "
            "Tokens present in the exported vocabulary accumulate raw counts, multiply by fitted "
            "IDF, L2-normalize, and form a decision score w\u00b7x + b. For this binary model the "
            "positive class is REAL, so P(REAL) = sigmoid(score) and P(FAKE) = 1 \u2212 P(REAL). The "
            "predicted label is the class with higher probability; confidence is that probability."
        ),
        size=12, align="justify", space_after=6, line_spacing=1.5,
    )

    # ---------- 3. IMPLEMENTATION ----------
    section_heading(doc, "3. Implementation")
    add_para(
        doc,
        (
            "Training lives in ml/train.py. The vectorizer uses lowercase=True, English stop "
            "words, max_features=5000, unigrams, and L2-normalized TF-IDF fit only on the "
            "training split (80%). LogisticRegression uses solver=liblinear. After evaluation, "
            "the script exports terms, idf, coef, intercept, and positiveLabel for the web app."
        ),
        size=12, align="justify", space_after=8, line_spacing=1.4,
    )
    add_para(doc, "Key training and export fragment (Python):", size=12, bold=True, align="left", space_after=4)

    code_snippet = r'''vectorizer = TfidfVectorizer(
    lowercase=True, stop_words="english",
    max_features=5000, ngram_range=(1, 1),
    token_pattern=r"(?u)\b\w\w+\b", norm="l2",
)
classifier = LogisticRegression(solver="liblinear", max_iter=1000,
                                random_state=42)
pipeline = Pipeline([("tfidf", vectorizer), ("clf", classifier)])
pipeline.fit(x_train, y_train)

# Export for JS: vocab order, idf, coef, intercept, positiveLabel
positive = str(clf.classes_[1])  # REAL for this binary fit
model_json = {
    "terms": terms, "idf": list(tfidf.idf_),
    "coef": list(clf.coef_.ravel()),
    "intercept": float(clf.intercept_.ravel()[0]),
    "positiveLabel": positive,
}'''
    add_code_block(doc, code_snippet, size=8)

    add_para(
        doc,
        (
            "Browser-side scoring in web/lib/predict.ts mirrors the sklearn pipeline closely on "
            "English ASCII text. The API route accepts POST { \"text\": \"...\" } and returns "
            "label, confidence, and probabilities. GET /api/predict returns the stored hold-out "
            "metrics for inspection."
        ),
        size=12, align="justify", space_before=8, space_after=6, line_spacing=1.4,
    )
    add_para(doc, "Key TypeScript inference fragment:", size=12, bold=True, align="left", space_after=4)

    ts_snippet = r'''export function predict(text: string): Prediction {
  const tokens = tokenize(text);  // /\b[a-z0-9_]{2,}\b/g
  // raw counts -> * idf -> L2 normalize -> decision = w·x + b
  const pPositive = sigmoid(decision);
  // positiveLabel === "REAL"
  const probabilities = {
    REAL: pPositive,
    FAKE: 1 - pPositive,
  };
  const label = probabilities.FAKE >= probabilities.REAL ? "FAKE" : "REAL";
  return { label, confidence: probabilities[label], probabilities };
}'''
    add_code_block(doc, ts_snippet, size=8)

    add_para(
        doc,
        (
            "Project layout: data/ holds the CSV (downloaded by train.py; gitignored), ml/ has "
            "train.py and optional predict.py, models/ stores metrics.json and the sklearn "
            "joblib artifact, and web/ is the Next.js app (App Router page, API route, "
            "lib/model.json, lib/predict.ts, lib/samples.json)."
        ),
        size=12, align="justify", space_before=8, space_after=6, line_spacing=1.4,
    )

    # ---------- TECHNOLOGIES USED ----------
    section_heading(doc, "Technologies Used")
    for name, role in [
        ("Python 3 + scikit-learn", "TF-IDF vectorizer, logistic regression, metrics, and model export."),
        ("pandas / NumPy / joblib", "Dataset loading, numeric arrays, and local sklearn pipeline persistence."),
        ("Next.js 15 (App Router)", "Web UI and /api/predict route for the laboratory demo."),
        ("TypeScript / React", "Client page and pure-JS TF-IDF + sigmoid inference."),
        ("McIntire fake_or_real_news", "Public labelled corpus (~6.3k articles) used for train/test."),
        ("Node.js", "Runtime for npm scripts and the Next.js development server."),
    ]:
        add_mixed_para(
            doc, [(f"\u2022  {name} \u2014 ", True), (role, False)],
            align="justify", size=12, space_after=8, line_spacing=1.4,
        )

    # ---------- 4. RESULTS ----------
    section_heading(doc, "4. Results")
    add_para(
        doc,
        (
            "The classifier was evaluated on a stratified 20% hold-out (train size 5063, test "
            "size 1266, random_state=42). Reported metrics from models/metrics.json are: "
            "accuracy 0.9202, F1 macro 0.9202, F1(FAKE) 0.9213, F1(REAL) 0.9191. The confusion "
            "matrix (rows = true label, columns = predicted FAKE, REAL) is [[591, 41], [60, 574]]."
        ),
        size=11, align="justify", space_after=4, line_spacing=1.2,
    )
    add_para(
        doc,
        (
            "The Next.js demo loads hold-out examples via Load fake example and Load real "
            "example. Predicting the fake sample yields FAKE at about 94% confidence; the real "
            "sample yields REAL at about 95% confidence, matching the stored hold-out "
            "probabilities. GET /api/predict returns the same metrics JSON used in this report."
        ),
        size=11, align="justify", space_after=4, line_spacing=1.2,
    )

    figures = [
        ("fig1-home-empty.png", "Fig. 1: Home screen \u2014 empty article input"),
        ("fig2-predict-fake.png", "Fig. 2: Prediction result \u2014 FAKE (~94% confidence)"),
        ("fig3-predict-real.png", "Fig. 3: Prediction result \u2014 REAL (~95% confidence)"),
    ]

    # Stack screenshots vertically so they print large enough to read.
    for fig_path, caption in figures:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(str(FIGS / fig_path), width=Inches(5.8))

        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(4)
        cr = cp.add_run(caption)
        cr.italic = True
        cr.font.size = Pt(10)
        cr.font.name = "Liberation Serif"

    add_para(
        doc,
        (
            "Summary: hold-out accuracy/F1 \u2248 0.92; UI screenshots confirm empty home, FAKE, and "
            "REAL prediction paths used for laboratory evaluation."
        ),
        size=11, align="justify", space_before=6, space_after=4, line_spacing=1.25,
    )

    # ---------- 5. CONCLUSION ----------
    section_heading(doc, "5. Conclusion")
    add_para(
        doc,
        (
            "This mini-project delivered a working Fake News Detection demo suitable for an NLP "
            "laboratory demonstration. A TF-IDF + logistic regression pipeline reaches about "
            "0.92 accuracy and macro F1 on a held-out portion of the McIntire corpus. Exporting "
            "the model to JSON enables TypeScript inference inside a Next.js app, so the demo "
            "runs without a Python server at request time."
        ),
        size=12, align="justify", space_after=10, line_spacing=1.5,
    )
    add_para(
        doc,
        (
            "The design deliberately stays small: classical features only, a single binary "
            "domain corpus, and no claim of general fact-checking. Future work may include "
            "n-gram or character features, calibrated probabilities, transformer baselines, "
            "broader and more recent news datasets, or explainability (top contributing terms). "
            "Within the stated scope, train \u2192 export \u2192 predict in the browser is implemented, "
            "measured, and documented."
        ),
        size=12, align="justify", space_after=6, line_spacing=1.5,
    )

    # ---------- 6. REFERENCES ----------
    section_heading(doc, "6. References")
    refs = [
        "G. Forman and others, classical TF-IDF and linear classifiers for text categorization \u2014 foundational NLP/ML practice used in this lab baseline.",
        "scikit-learn Developers, \u201cscikit-learn Documentation \u2014 TfidfVectorizer & LogisticRegression,\u201d https://scikit-learn.org/stable/ (accessed 2026).",
        "McIntire, G. / Lutz Hamel mirror, \u201cfake_or_real_news.csv,\u201d https://github.com/lutzhamel/fake-news (dataset used for training).",
        "Vercel, \u201cNext.js Documentation,\u201d https://nextjs.org/docs (accessed 2026).",
        "Pedregosa et al., \u201cScikit-learn: Machine Learning in Python,\u201d JMLR, 2011.",
        "Shree L. R. Tiwari College Of Engineering, Department of Computer Engineering \u2014 NLP Mini Project guidelines (internal).",
    ]
    for i, r in enumerate(refs, 1):
        add_para(doc, f"{i}. {r}", size=12, align="justify", space_after=8, line_spacing=1.35)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))
    print(f"Wrote {OUT_DOCX}")


if __name__ == "__main__":
    build()
